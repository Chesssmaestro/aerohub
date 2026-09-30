"""Импорт новых запчастей из zip-каталогов DJI (T50/T55/T70P/T100 Shopping Cart).

Разово запускается вручную, дописывает новые позиции в data/parts.json
(не трогает уже существующие 2544 позиции). Источник — распакованные
"... Catalog.xlsx" / "T55.xlsx", лист "Shopping Cart": Image/Items/Material Number/Product Line.

Названия оставлены на английском (решение пользователя): в отличие от остального
каталога, у этих позиций нет цены/наличия с сайта-источника — ставится
price=0, stock='Под заказ', группа/тип определяются по английским ключевым словам.

Ожидает рядом с репозиторием (корень проекта) исходные архивы DJI T100.zip,
DJI T50.zip, DJI T55.zip, DJI T70.zip — не хранятся в git (см. .gitignore).

Запуск:  .venv\\Scripts\\pip install openpyxl
         .venv\\Scripts\\python tools\\import_dji_catalogs.py
"""

import io
import json
import re
import zipfile
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
PARTS_JSON = BASE_DIR / 'data' / 'parts.json'

# zip -> имя xlsx с листом "Shopping Cart" внутри архива
CATALOGS = {
    'DJI T100.zip': 'DJI T100/T100 Catalog.xlsx',
    'DJI T50.zip': 'DJI T50/T50 Catalog.xlsx',
    'DJI T55.zip': 'DJI T55/T55.xlsx',
    'DJI T70.zip': 'DJI T70/T70P Catalog.xlsx',
}

MODEL_RE = re.compile(r'\bT\d{2,3}P?\b|\bJ150\b')

# Группа узла: ключевые слова на английском (источник — DJI-каталог без перевода),
# тот же набор групп, что и в tools/fetch_parts.py, чтобы фильтры сайта не дробились.
GROUPS = [
    ('Опрыскивание', ('nozzle', 'sprinkler', 'sprayer', 'spray lance', 'spray boom')),
    ('Насосы и расход', ('pump', 'flow meter', 'flowmeter', 'valve', 'hose', 'fitting', 'tube', 'tubing', 'pipe')),
    ('Бак и жидкостная система', ('spray tank', 'tank', 'filter', 'liquid level', 'level sensor')),
    ('Внесение гранул', ('granule', 'hopper', 'auger', 'feeder', 'spreader', 'seeder')),
    ('Лучи и моторы', ('aircraft arm', 'arm ', 'motor', 'esc', 'propulsion')),
    ('Винты и пропеллеры', ('propeller', 'blade', 'folding prop')),
    ('Шасси и рама', ('landing gear', 'frame', 'chassis', 'bracket', 'housing', 'base', 'stand', 'leg')),
    ('Аккумуляторы и питание', ('battery', 'charger', 'charging', 'power supply', 'generator', 'power cable')),
    ('Электроника и датчики', ('board', 'module', 'sensor', 'radar', 'camera', 'antenna', 'gnss', 'rtk', 'led')),
    ('Пульт управления', ('remote control', 'joystick', 'stick', 'controller screen')),
    ('Кабели и разъёмы', ('cable', 'connector', 'wire', 'plug', 'flex cable', 'wiring harness')),
    ('Крепёж и уплотнения', ('screw', 'bolt', 'nut', 'washer', 'rivet', 'seal', 'gasket', 'ring',
                             'o-ring', 'clamp', 'stopper', 'self-tapping', 'wrench')),
]

TYPES = [
    ('Крепёж', ('screw', 'bolt', 'nut', 'washer', 'rivet', 'pin', 'clamp', 'wrench')),
    ('Уплотнения', ('seal', 'gasket', 'o-ring', 'oring', 'sealing ring', 'ring')),
    ('Кабели и разъёмы', ('cable', 'connector', 'wire', 'flex cable', 'plug')),
    ('Платы и модули', ('board', 'module', 'controller', 'control unit', 'pcb')),
    ('Датчики', ('sensor', 'radar', 'camera')),
    ('Насосы', ('pump',)),
    ('Распылители', ('nozzle', 'sprinkler', 'spray head')),
    ('Пропеллеры', ('propeller', 'blade')),
    ('Моторы', ('motor',)),
    ('Корпусные детали', ('housing', 'cover', 'panel', 'shell', 'cap')),
    ('Баки и ёмкости', ('tank', 'hopper', 'container')),
    ('Шланги и фитинги', ('hose', 'fitting', 'tube', 'pipe', 'nipple')),
    ('Аккумуляторы', ('battery',)),
    ('Антенны', ('antenna',)),
]


def classify(name: str, table: list[tuple[str, tuple[str, ...]]], default: str) -> str:
    low = name.lower()
    for title, keys in table:
        if any(k in low for k in keys):
            return title
    return default


def compat_models(text: str) -> list[str]:
    return sorted(set(MODEL_RE.findall(text)), key=lambda m: (len(m), m))


def read_catalog(zip_path: Path, inner_name: str) -> dict[str, dict]:
    with zipfile.ZipFile(zip_path) as zf:
        data = zf.read(inner_name)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb['Shopping Cart']
    out: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
        _img, item, article, line = row
        if not article or not item:
            continue
        article = str(article).strip()
        item = str(item).strip()
        entry = out.setdefault(article, {'name': item, 'lines': []})
        entry['lines'].append(line or '')
    wb.close()
    return out


def main() -> None:
    by_article: dict[str, dict] = {}
    for zip_name, inner_name in CATALOGS.items():
        for article, entry in read_catalog(BASE_DIR / zip_name, inner_name).items():
            merged = by_article.setdefault(article, {'name': entry['name'], 'lines': []})
            merged['lines'].extend(entry['lines'])

    existing = json.loads(PARTS_JSON.read_text(encoding='utf-8'))
    existing_articles = {p['article'] for p in existing}
    next_id = max(p['external_id'] for p in existing) + 1

    added = []
    for article in sorted(by_article):
        if article in existing_articles:
            continue
        entry = by_article[article]
        name = entry['name']
        models = compat_models(' '.join(entry['lines']))
        added.append({
            'external_id': next_id,
            'name': name,
            'article': article,
            'price': 0,
            'stock': 'Под заказ',
            'group': classify(name, GROUPS, 'Прочие узлы'),
            'kind': classify(name, TYPES, 'Прочее'),
            'models': models,
        })
        next_id += 1

    combined = existing + added
    combined.sort(key=lambda x: x['external_id'])
    PARTS_JSON.write_text(json.dumps(combined, ensure_ascii=False, indent=1), encoding='utf-8')
    print(f'было {len(existing)}, добавлено {len(added)}, стало {len(combined)}')


if __name__ == '__main__':
    main()
